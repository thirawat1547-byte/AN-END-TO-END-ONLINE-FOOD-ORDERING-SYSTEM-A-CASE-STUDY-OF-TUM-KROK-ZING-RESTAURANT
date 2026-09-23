import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Inches
import sys

def format_cell_text(cell, text, bold=False, font_name="TH SarabunPSK", size_pt=16):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = 0
    p.paragraph_format.space_after = 0
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    if bold:
        b = OxmlElement('w:b')
        bCs = OxmlElement('w:bCs')
        rPr.append(b)
        rPr.append(bCs)

def set_table_row(row, data, bold=False):
    for i, val in enumerate(data):
        if i < len(row.cells):
            format_cell_text(row.cells[i], val, bold=bold)

def make_heading_paragraph(doc, text, font_name="TH SarabunPSK", size_pt=16, bold=True, space_before=Pt(6), space_after=Pt(4)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    if bold:
        rPr.append(OxmlElement('w:b'))
        rPr.append(OxmlElement('w:bCs'))
    return p

def make_body_paragraph(doc, text, font_name="TH SarabunPSK", size_pt=16, space_before=0, space_after=Pt(4)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    return p

def make_styled_table(doc, headers, rows_data):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for c_i, h in enumerate(headers):
        format_cell_text(table.rows[0].cells[c_i], h, bold=True)
        shading_elm = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
        table.rows[0].cells[c_i]._tc.get_or_add_tcPr().append(shading_elm)
    for r_i, r_data in enumerate(rows_data):
        row = table.rows[r_i + 1]
        set_table_row(row, r_data, bold=False)
    return table

def insert_after(current_element, new_element):
    current_element.addnext(new_element)
    return new_element

def update_main_thesis(thesis_path):
    print(f"Loading main thesis: {thesis_path}...")
    doc = docx.Document(thesis_path)

    # 1. Update UC-01 (Table 12)
    t_uc01 = None
    for i, t in enumerate(doc.tables):
        if len(t.rows) > 1 and len(t.rows[1].cells) >= 2:
            if t.rows[1].cells[0].text.strip() == 'รหัสยูสเคส' and t.rows[1].cells[1].text.strip() == 'UC-01':
                t_uc01 = t
                print(f"Found UC-01 at table index {i}")
                break

    if t_uc01:
        print("Updating UC-01 table...")
        uc01_data = [
            ["Use Case ID", "UC-01"],
            ["Use Case Name", "เข้าสู่ระบบและยืนยันตัวตน (User Authentication)"],
            ["Actors", "ลูกค้า (Customer), พนักงานห้องครัว (Kitchen Staff), ผู้ดูแลระบบ (Admin), พนักงานส่งอาหาร (Rider)"],
            ["Pre-condition", "ผู้ใช้งานเปิดหน้าเว็บแอปพลิเคชันและเข้าสู่หน้าจอเข้าสู่ระบบ (/login)"],
            ["Post-condition", "ระบบออกโทเคนยืนยันสิทธิ์ (JWT Access Token) และนำผู้ใช้ไปยังหน้าจอตามสิทธิ์ (Role)"],
            ["กระบวนการทำงานหลัก (Main Flow)", "1. ผู้ใช้งานกรอกชื่อผู้ใช้ (Username) หรือเบอร์โทรศัพท์ และรหัสผ่าน (Password)\n2. ผู้ใช้งานกดปุ่ม 'เข้าสู่ระบบ'\n3. ระบบทำการตรวจสอบความถูกต้องของข้อมูลผ่านระบบความปลอดภัย (Bcrypt Hash Verification)\n4. เมื่อข้อมูลถูกต้อง ระบบสร้าง JSON Web Token (JWT) และจัดเก็บสถานะเข้าสู่ระบบไว้ใน Client Storage\n5. ระบบนำทาง (Redirect) ผู้ใช้งานไปยังหน้าแรกตามระดับสิทธิ์:\n   - สิทธิ์ ADMIN -> หน้าจอจัดการร้าน (/admin)\n   - สิทธิ์ KITCHEN -> หน้าจอห้องครัว KDS (/kitchen)\n   - สิทธิ์ RIDER -> หน้าจอรับงานจัดส่ง (/rider)\n   - สิทธิ์ CUSTOMER -> หน้าจอสั่งอาหาร (/menu หรือ /)"],
            ["กระบวนการทำงานทางเลือก (Alternative Flows)", "3a. ข้อมูลเข้าสู่ระบบไม่ถูกต้อง:\n    1. ระบบตรวจสอบพบว่าชื่อผู้ใช้ไม่มีอยู่ในระบบ หรือรหัสผ่านไม่ถูกต้อง\n    2. ระบบแสดงข้อความแจ้งเตือนสีแดง 'ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง กรุณาลองใหม่อีกครั้ง'\n    3. ผู้ใช้คงอยู่ที่หน้าเดิมเพื่อแก้ไขข้อมูล\n3b. ยังไม่มีบัญชีผู้ใช้งาน:\n    1. ผู้ใช้งานคลิกลิงก์ 'สมัครสมาชิกใหม่'\n    2. ระบบเปลี่ยนหน้าจอไปยังหน้าลงทะเบียน (/register)"]
        ]
        # set header
        format_cell_text(t_uc01.rows[0].cells[0], "หัวข้อ", bold=True)
        format_cell_text(t_uc01.rows[0].cells[1], "รายละเอียด", bold=True)
        shd = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
        t_uc01.rows[0].cells[0]._tc.get_or_add_tcPr().append(shd)
        shd2 = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
        t_uc01.rows[0].cells[1]._tc.get_or_add_tcPr().append(shd2)

        while len(t_uc01.rows) > len(uc01_data) + 1:
            t_uc01._tbl.remove(t_uc01.rows[-1]._tr)
        while len(t_uc01.rows) < len(uc01_data) + 1:
            t_uc01.add_row()
        for i, row_data in enumerate(uc01_data):
            set_table_row(t_uc01.rows[i + 1], row_data)

    # 2. Insert UC-16 (After Table 26 UC-15)
    t_uc15 = None
    for i, t in enumerate(doc.tables):
        if len(t.rows) > 1 and len(t.rows[1].cells) >= 2:
            if t.rows[1].cells[0].text.strip() == 'รหัสยูสเคส' and t.rows[1].cells[1].text.strip() == 'UC-15':
                t_uc15 = t
                print(f"Found UC-15 at table index {i}")
                break

    if t_uc15:
        # Check if UC-16 already added
        already_has_uc16 = any("UC-16" in p.text for p in doc.paragraphs)
        if not already_has_uc16:
            print("Inserting UC-16 table after UC-15...")
            cur = t_uc15._element
            p_uc16_space = make_body_paragraph(doc, "", space_after=Pt(4))
            cur = insert_after(cur, p_uc16_space._element)

            p_uc16_h = make_heading_paragraph(doc, "ตารางที่ 3.16 คำอธิบายยูสเคส รหัส UC-16", bold=True, space_before=Pt(8), space_after=Pt(4))
            cur = insert_after(cur, p_uc16_h._element)

            uc16_headers = ["หัวข้อ", "รายละเอียด"]
            uc16_data = [
                ["Use Case ID", "UC-16"],
                ["Use Case Name", "จัดการการส่งอาหารและตรวจสอบสถานะการจัดส่ง (Rider Order Delivery)"],
                ["Actors", "พนักงานส่งอาหาร (Rider)"],
                ["Pre-condition", "ไรเดอร์เข้าสู่ระบบด้วยสิทธิ์ RIDER และมีออเดอร์เดลิเวอรีที่ครัวทำเสร็จแล้ว (สถานะ READY)"],
                ["Post-condition", "สถานะออเดอร์เปลี่ยนเป็น 'จัดส่งเรียบร้อย' (DELIVERED) พร้อมบันทึกประวัติการส่งสำเร็จ"],
                ["กระบวนการทำงานหลัก (Main Flow)", "1. ไรเดอร์เข้าสู่หน้าจอ 'ระบบรับงานและจัดส่งอาหาร' (/rider)\n2. ระบบแสดงรายการออเดอร์เดลิเวอรีที่มีสถานะพร้อมส่ง (READY) พร้อมรายละเอียดที่อยู่จัดส่งและเบอร์ติดต่อลูกค้า\n3. ไรเดอร์กดปุ่ม 'รับงานจัดส่ง' (Accept Job)\n4. ระบบปรับสถานะออเดอร์เป็น 'กำลังจัดส่ง' (IN_DELIVERY)\n5. ไรเดอร์เดินทางไปส่งอาหารยังที่อยู่ของลูกค้า พร้อมสามารถกดโทรติดต่อลูกค้าได้โดยตรง\n6. เมื่อส่งมอบอาหารและรับชำระเงินเรียบร้อย ไรเดอร์กดปุ่ม 'ส่งสำเร็จ' (Complete Delivery)\n7. ระบบปรับสถานะออเดอร์เป็น 'จัดส่งเรียบร้อย' (DELIVERED) และบันทึกประวัติลงฐานข้อมูล"],
                ["กระบวนการทำงานทางเลือก (Alternative Flows)", "2a. ไม่พบงานที่พร้อมส่งในระบบ:\n    1. หน้าจอแสดงข้อความ 'ไม่มีรายการอาหารที่พร้อมจัดส่งในขณะนี้'\n    2. ระบบทำการตรวจสอบคิวใหม่อัตโนมัติ (Auto Refresh) ทุกๆ ช่วงเวลาที่กำหนด"]
            ]
            t_uc16 = make_styled_table(doc, uc16_headers, uc16_data)
            cur = insert_after(cur, t_uc16._element)

    # 3. Update Real-time (WebSocket -> HTTP Polling) in text paragraphs and tables
    print("Updating Real-time (WebSocket -> HTTP Polling) text...")
    polling_text = "ในส่วนของการติดตามสถานะออเดอร์แบบทันท่วงที (Real-time Order Tracking) ระบบได้รับการออกแบบโดยใช้สถาปัตยกรรม Lightweight HTTP Polling Architecture ผสานการทำงานร่วมกับ Vue Reactive State Management (Pinia) โดยหน้าจอแสดงผลของลูกค้า (Tracking View) และหน้าจอห้องครัว (Kitchen Display System - KDS) จะมีการส่งคำขอแบบ Asynchronous HTTP Request เพื่อซิงโครไนซ์สถานะคำสั่งซื้อจาก RESTful API อย่างต่อเนื่องตามรอบเวลาที่กำหนด (Configurable Polling Interval) ซึ่งช่วยลด Overhead ในการรักษา TCP Connection ตลอดเวลาของฝั่งเซิร์ฟเวอร์ ให้ความเสถียรสูง รองรับการทำงานผ่าน Reverse Proxy และ Cloud Infrastructure ได้อย่างมีประสิทธิภาพ ปราศจากปัญหา Connection Dropped ของ Stateful WebSocket"

    # Search and replace in paragraphs
    for p in doc.paragraphs:
        txt = p.text
        if "Real-time Communication: WebSocket" in txt:
            p.text = "Real-time Communication: สถาปัตยกรรม Lightweight HTTP Polling ผสาน Vue Reactive State (Pinia) สำหรับซิงโครไนซ์ข้อมูลและอัปเดตสถานะของคำสั่งซื้อแบบเรียลไทม์อย่างมีเสถียรภาพ"
            for r in p.runs:
                r.font.name = "TH SarabunPSK"
        elif "Native Support for WebSockets" in txt:
            p.text = "High Performance RESTful API: มี Module และ Decorators สำหรับจัดการ Request/Response และรองรับการทำ Data Synchronization และ Polling ได้อย่างมีประสิทธิภาพและเสถียรภาพสูง"
            for r in p.runs:
                r.font.name = "TH SarabunPSK"
        elif "2.7.1 เทคโนโลยี WebSockets (WSS Protocol)" in txt:
            p.text = "2.7.1 สถาปัตยกรรมการสื่อสารข้อมูลทันท่วงที (Lightweight HTTP Polling Architecture)"
            for r in p.runs:
                r.font.name = "TH SarabunPSK"
        elif "WebSockets เป็นโพรโทคอลการสื่อสารข้อมูลแบบสองทาง" in txt:
            p.text = polling_text
            for r in p.runs:
                r.font.name = "TH SarabunPSK"
        elif "การแจ้งเตือนเรียลไทม์ไปยังห้องครัวและการปรุงอาหาร" in txt and "WebSockets" in txt:
            p.text = p.text.replace("ผ่าน WebSockets (WSS)", "ผ่านระบบ Asynchronous Polling และ RESTful API")
            for r in p.runs:
                r.font.name = "TH SarabunPSK"
        elif "และควบคุม WebSockets Gateway" in txt:
            p.text = p.text.replace("และควบคุม WebSockets Gateway", "และให้บริการ RESTful API พร้อมกลไก Data Synchronization สำหรับการติดตามสถานะออเดอร์")
            for r in p.runs:
                r.font.name = "TH SarabunPSK"
        elif "และรองรับการเชื่อมต่อแบบ WebSockets สำหรับการแจ้งเตือนสถานะการปรุงอาหารและสถานะโต๊ะแบบเรียลไทม์" in txt:
            p.text = p.text.replace("และรองรับการเชื่อมต่อแบบ WebSockets สำหรับการแจ้งเตือนสถานะการปรุงอาหารและสถานะโต๊ะแบบเรียลไทม์", "และรองรับการซิงโครไนซ์ข้อมูลแบบทันท่วงที (Real-time Polling Synchronization) สำหรับการแจ้งเตือนสถานะการปรุงอาหารและสถานะโต๊ะ")
            for r in p.runs:
                r.font.name = "TH SarabunPSK"

    # Search and replace in tables
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if "ระบบต้องอัปเดตสถานะออเดอร์ไปยังจอ KDS แบบเรียลไทม์ผ่าน WebSockets" in c.text:
                    format_cell_text(c, "ระบบต้องอัปเดตสถานะออเดอร์ไปยังจอ KDS แบบทันท่วงที (Real-time / Polling Synchronization)")
                elif "ความหน่วง (Latency) ของการส่งข้อมูลผ่าน WebSockets" in c.text:
                    format_cell_text(c, "ความหน่วง (Latency) ของการส่งข้อมูลและการตอบสนองของระบบ Polling")

    doc.save(thesis_path)
    print(f"Saved main thesis: {thesis_path} successfully!")

def update_datadict_file(datadict_path):
    print(f"Updating datadict file: {datadict_path}...")
    doc = docx.Document(datadict_path)

    # Check if Part 3, 4, 5 already added
    has_part3 = any("ส่วนที่ 3" in p.text for p in doc.paragraphs)
    if not has_part3:
        # Part 3: Use Case UC-01
        p_p3_h = doc.add_paragraph()
        p_p3_h.paragraph_format.space_before = Pt(14)
        p_p3_h.paragraph_format.space_after = Pt(4)
        r = p_p3_h.add_run("ส่วนที่ 3: Use Case UC-01 (เข้าสู่ระบบและยืนยันตัวตน) ฉบับปรับปรุง")
        r.font.name = "TH SarabunPSK"
        r.font.size = Pt(16)
        r.bold = True

        p_p3_d = doc.add_paragraph()
        p_p3_d.paragraph_format.space_after = Pt(6)
        r = p_p3_d.add_run("นำไปแทนที่ Use Case UC-01 ใน บทที่ 3 หัวข้อ 3.3 โดยตัด Social Login ออกเพื่อให้สอดคล้องกับระบบยืนยันตัวตนจริงของร้าน")
        r.font.name = "TH SarabunPSK"
        r.font.size = Pt(16)

        uc01_headers = ["หัวข้อ", "รายละเอียด"]
        uc01_data = [
            ["Use Case ID", "UC-01"],
            ["Use Case Name", "เข้าสู่ระบบและยืนยันตัวตน (User Authentication)"],
            ["Actors", "ลูกค้า (Customer), พนักงานห้องครัว (Kitchen Staff), ผู้ดูแลระบบ (Admin), พนักงานส่งอาหาร (Rider)"],
            ["Pre-condition", "ผู้ใช้งานเปิดหน้าเว็บแอปพลิเคชันและเข้าสู่หน้าจอเข้าสู่ระบบ (/login)"],
            ["Post-condition", "ระบบออกโทเคนยืนยันสิทธิ์ (JWT Access Token) และนำผู้ใช้ไปยังหน้าจอตามสิทธิ์ (Role)"],
            ["กระบวนการทำงานหลัก (Main Flow)", "1. ผู้ใช้งานกรอกชื่อผู้ใช้ (Username) หรือเบอร์โทรศัพท์ และรหัสผ่าน (Password)\n2. ผู้ใช้งานกดปุ่ม 'เข้าสู่ระบบ'\n3. ระบบทำการตรวจสอบความถูกต้องของข้อมูลผ่านระบบความปลอดภัย (Bcrypt Hash Verification)\n4. เมื่อข้อมูลถูกต้อง ระบบสร้าง JSON Web Token (JWT) และจัดเก็บสถานะเข้าสู่ระบบไว้ใน Client Storage\n5. ระบบนำทาง (Redirect) ผู้ใช้งานไปยังหน้าแรกตามระดับสิทธิ์:\n   - สิทธิ์ ADMIN -> หน้าจอจัดการร้าน (/admin)\n   - สิทธิ์ KITCHEN -> หน้าจอห้องครัว KDS (/kitchen)\n   - สิทธิ์ RIDER -> หน้าจอรับงานจัดส่ง (/rider)\n   - สิทธิ์ CUSTOMER -> หน้าจอสั่งอาหาร (/menu หรือ /)"],
            ["กระบวนการทำงานทางเลือก (Alternative Flows)", "3a. ข้อมูลเข้าสู่ระบบไม่ถูกต้อง:\n    1. ระบบตรวจสอบพบว่าชื่อผู้ใช้ไม่มีอยู่ในระบบ หรือรหัสผ่านไม่ถูกต้อง\n    2. ระบบแสดงข้อความแจ้งเตือนสีแดง 'ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง กรุณาลองใหม่อีกครั้ง'\n    3. ผู้ใช้คงอยู่ที่หน้าเดิมเพื่อแก้ไขข้อมูล\n3b. ยังไม่มีบัญชีผู้ใช้งาน:\n    1. ผู้ใช้งานคลิกลิงก์ 'สมัครสมาชิกใหม่'\n    2. ระบบเปลี่ยนหน้าจอไปยังหน้าลงทะเบียน (/register)"]
        ]
        make_styled_table(doc, uc01_headers, uc01_data)

        # Part 4: Use Case UC-16
        p_p4_h = doc.add_paragraph()
        p_p4_h.paragraph_format.space_before = Pt(14)
        p_p4_h.paragraph_format.space_after = Pt(4)
        r = p_p4_h.add_run("ส่วนที่ 4: Use Case UC-16 ระบบรับงานส่งอาหารโดยไรเดอร์ (Rider Delivery)")
        r.font.name = "TH SarabunPSK"
        r.font.size = Pt(16)
        r.bold = True

        p_p4_d = doc.add_paragraph()
        p_p4_d.paragraph_format.space_after = Pt(6)
        r = p_p4_d.add_run("เพิ่มเป็น Use Case ใหม่ใน บทที่ 3 หัวข้อ 3.3 เพื่อสะท้อนถึงระบบสั่งอาหารแบบครบวงจรที่มีระบบจัดส่งอาหารรองรับ")
        r.font.name = "TH SarabunPSK"
        r.font.size = Pt(16)

        uc16_headers = ["หัวข้อ", "รายละเอียด"]
        uc16_data = [
            ["Use Case ID", "UC-16"],
            ["Use Case Name", "จัดการการส่งอาหารและตรวจสอบสถานะการจัดส่ง (Rider Order Delivery)"],
            ["Actors", "พนักงานส่งอาหาร (Rider)"],
            ["Pre-condition", "ไรเดอร์เข้าสู่ระบบด้วยสิทธิ์ RIDER และมีออเดอร์เดลิเวอรีที่ครัวทำเสร็จแล้ว (สถานะ READY)"],
            ["Post-condition", "สถานะออเดอร์เปลี่ยนเป็น 'จัดส่งเรียบร้อย' (DELIVERED) พร้อมบันทึกประวัติการส่งสำเร็จ"],
            ["กระบวนการทำงานหลัก (Main Flow)", "1. ไรเดอร์เข้าสู่หน้าจอ 'ระบบรับงานและจัดส่งอาหาร' (/rider)\n2. ระบบแสดงรายการออเดอร์เดลิเวอรีที่มีสถานะพร้อมส่ง (READY) พร้อมรายละเอียดที่อยู่จัดส่งและเบอร์ติดต่อลูกค้า\n3. ไรเดอร์กดปุ่ม 'รับงานจัดส่ง' (Accept Job)\n4. ระบบปรับสถานะออเดอร์เป็น 'กำลังจัดส่ง' (IN_DELIVERY)\n5. ไรเดอร์เดินทางไปส่งอาหารยังที่อยู่ของลูกค้า พร้อมสามารถกดโทรติดต่อลูกค้าได้โดยตรง\n6. เมื่อส่งมอบอาหารและรับชำระเงินเรียบร้อย ไรเดอร์กดปุ่ม 'ส่งสำเร็จ' (Complete Delivery)\n7. ระบบปรับสถานะออเดอร์เป็น 'จัดส่งเรียบร้อย' (DELIVERED) และบันทึกประวัติลงฐานข้อมูล"],
            ["กระบวนการทำงานทางเลือก (Alternative Flows)", "2a. ไม่พบงานที่พร้อมส่งในระบบ:\n    1. หน้าจอแสดงข้อความ 'ไม่มีรายการอาหารที่พร้อมจัดส่งในขณะนี้'\n    2. ระบบทำการตรวจสอบคิวใหม่อัตโนมัติ (Auto Refresh) ทุกๆ ช่วงเวลาที่กำหนด"]
        ]
        make_styled_table(doc, uc16_headers, uc16_data)

        # Part 5: Real-time Polling Text
        p_p5_h = doc.add_paragraph()
        p_p5_h.paragraph_format.space_before = Pt(14)
        p_p5_h.paragraph_format.space_after = Pt(4)
        r = p_p5_h.add_run("ส่วนที่ 5: ข้อความวิชาการปรับแก้เรื่อง Real-time (WebSocket -> HTTP Polling)")
        r.font.name = "TH SarabunPSK"
        r.font.size = Pt(16)
        r.bold = True

        p_p5_d = doc.add_paragraph()
        p_p5_d.paragraph_format.space_after = Pt(6)
        r = p_p5_d.add_run("สำหรับนำไปแทนที่ข้อความใน บทที่ 2 และ บทที่ 3 ในส่วนที่เคยกล่าวถึง WebSocket เพื่อให้ตรงกับสถาปัตยกรรมระบบจริง:")
        r.font.name = "TH SarabunPSK"
        r.font.size = Pt(16)

        p_p5_text = doc.add_paragraph()
        p_p5_text.paragraph_format.space_after = Pt(8)
        p_p5_text.paragraph_format.line_spacing = 1.15
        r = p_p5_text.add_run('\"ในส่วนของการติดตามสถานะออเดอร์แบบทันท่วงที (Real-time Order Tracking) ระบบได้รับการออกแบบโดยใช้สถาปัตยกรรม Lightweight HTTP Polling Architecture ผสานการทำงานร่วมกับ Vue Reactive State Management (Pinia) โดยหน้าจอแสดงผลของลูกค้า (Tracking View) และหน้าจอห้องครัว (Kitchen Display System - KDS) จะมีการส่งคำขอแบบ Asynchronous HTTP Request เพื่อซิงโครไนซ์สถานะคำสั่งซื้อจาก RESTful API อย่างต่อเนื่องตามรอบเวลาที่กำหนด (Configurable Polling Interval) ซึ่งช่วยลด Overhead ในการรักษา TCP Connection ตลอดเวลาของฝั่งเซิร์ฟเวอร์ ให้ความเสถียรสูง รองรับการทำงานผ่าน Reverse Proxy และ Cloud Infrastructure ได้อย่างมีประสิทธิภาพ ปราศจากปัญหา Connection Dropped ของ Stateful WebSocket\"')
        r.font.name = "TH SarabunPSK"
        r.font.size = Pt(16)

    doc.save(datadict_path)
    print(f"Saved datadict file: {datadict_path} successfully!")

if __name__ == "__main__":
    thesis_file = "ระบบสั่งอาหารออนไลน์แบบครบวงจร-กรณีศึกษาร้านตำครกซิ่ง-ล่าสุด.docx"
    datadict_file = "แก้ไขdatadictionary.docx"
    
    update_main_thesis(thesis_file)
    update_datadict_file(datadict_file)
