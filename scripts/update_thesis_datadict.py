import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Inches
import sys

def format_cell_text(cell, text, bold=False, font_name="TH SarabunPSK", size_pt=16):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = 0
    p.paragraph_format.space_after = 0
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    if bold:
        b = OxmlElement('w:b')
        bCs = OxmlElement('w:bCs')
        rPr.append(b)
        rPr.append(bCs)

def set_table_row(row, data, bold=False):
    for i, val in enumerate(data):
        if i < len(row.cells):
            format_cell_text(row.cells[i], val, bold=bold)

def make_heading_paragraph(doc, text, font_name="TH SarabunPSK", size_pt=16, bold=True, space_before=Pt(6), space_after=Pt(4)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    if bold:
        rPr.append(OxmlElement('w:b'))
        rPr.append(OxmlElement('w:bCs'))
    return p

def make_body_paragraph(doc, text, font_name="TH SarabunPSK", size_pt=16, space_before=0, space_after=Pt(4)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    return p

def make_styled_table(doc, headers, rows_data):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    
    # header
    for c_i, h in enumerate(headers):
        format_cell_text(table.rows[0].cells[c_i], h, bold=True)
        shading_elm = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
        table.rows[0].cells[c_i]._tc.get_or_add_tcPr().append(shading_elm)

    # rows
    for r_i, r_data in enumerate(rows_data):
        row = table.rows[r_i + 1]
        set_table_row(row, r_data, bold=False)

    return table

def insert_after(current_element, new_element):
    current_element.addnext(new_element)
    return new_element

def update_document(doc_path):
    doc = docx.Document(doc_path)
    print("Opened document successfully.")

    # Find the target tables dynamically
    t_users = None
    t_promo = None
    t_order_items = None
    t_ingredients = None
    t_menu_ingredients = None

    for i, t in enumerate(doc.tables):
        if len(t.rows) > 2 and len(t.rows[1].cells) >= 4:
            c1 = t.rows[1].cells[0].text.strip()
            c2 = t.rows[2].cells[0].text.strip()
            if c1 == 'user_id' and c2 == 'username':
                t_users = t
                print(f"Found USERS at table index {i}")
            elif c1 == 'promo_id' and c2 == 'code':
                t_promo = t
                print(f"Found PROMOTIONS at table index {i}")
            elif c1 == 'order_item_id' and c2 == 'order_id':
                t_order_items = t
                print(f"Found ORDER_ITEMS at table index {i}")
            elif c1 == 'ingredient_id' and ('name' in c2.lower() or 'ingredient' in c2.lower()):
                t_ingredients = t
                print(f"Found INGREDIENTS at table index {i}")
            elif c1 == 'menu_id' and c2 == 'ingredient_id':
                t_menu_ingredients = t
                print(f"Found MENU_INGREDIENTS at table index {i}")

    # 1. Update USERS table
    if t_users:
        print("Updating USERS table...")
        clean_users_data = [
            ["user_id", "INT", "PK", "รหัสอ้างอิงผู้ใช้งาน (เพิ่มค่าอัตโนมัติ)"],
            ["username", "VARCHAR(100)", "-", "ชื่อผู้ใช้งาน / ชื่อบัญชี (กำหนดค่าไม่ซ้ำ Unique)"],
            ["password", "VARCHAR(255)", "-", "รหัสผ่าน (เข้ารหัสความปลอดภัย bcrypt)"],
            ["email", "VARCHAR(150)", "-", "อีเมลสำหรับติดต่อและเข้าสู่ระบบ"],
            ["phone_number", "VARCHAR(15)", "-", "เบอร์โทรศัพท์สำหรับติดต่อจัดส่ง"],
            ["address", "VARCHAR(255)", "-", "ที่อยู่สำหรับติดต่อจัดส่งอาหาร (รองรับคำสั่งซื้อ Delivery)"],
            ["role", "VARCHAR(20)", "-", "สิทธิ์การใช้งาน (เช่น Customer, Staff, Kitchen, Admin)"]
        ]
        while len(t_users.rows) > len(clean_users_data) + 1:
            t_users._tbl.remove(t_users.rows[-1]._tr)
        while len(t_users.rows) < len(clean_users_data) + 1:
            t_users.add_row()
        for i, row_data in enumerate(clean_users_data):
            set_table_row(t_users.rows[i + 1], row_data)

    # 2. Update PROMOTIONS table
    if t_promo:
        print("Updating PROMOTIONS table...")
        clean_promo_data = [
            ["promo_id", "INT", "PK", "รหัสอ้างอิงโปรโมชัน (เพิ่มค่าอัตโนมัติ)"],
            ["code", "VARCHAR(50)", "-", "รหัสคูปองส่วนลดที่ลูกค้าต้องกรอก (กำหนดค่าไม่ซ้ำ Unique)"],
            ["discount_type", "VARCHAR(20)", "-", "ประเภทส่วนลด (เช่น Fixed, Percentage)"],
            ["discount_value", "DECIMAL(10,2)", "-", "มูลค่าส่วนลดที่ได้รับ"],
            ["min_order_price", "DECIMAL(10,2)", "-", "ยอดสั่งซื้อขั้นต่ำที่สามารถใช้โปรโมชันได้ (ค่าเริ่มต้น 0.00)"],
            ["expiry_date", "DATETIME", "-", "วันและเวลาที่โปรโมชันหมดอายุ"],
            ["created_at", "DATETIME", "-", "วันและเวลาที่สร้างรหัสโปรโมชัน (ค่าเริ่มต้นเวลาปัจจุบัน)"]
        ]
        while len(t_promo.rows) > len(clean_promo_data) + 1:
            t_promo._tbl.remove(t_promo.rows[-1]._tr)
        while len(t_promo.rows) < len(clean_promo_data) + 1:
            t_promo.add_row()
        for i, row_data in enumerate(clean_promo_data):
            set_table_row(t_promo.rows[i + 1], row_data)

    # 3. Update ORDER_ITEMS table
    if t_order_items:
        print("Updating ORDER_ITEMS table...")
        clean_order_items_data = [
            ["order_item_id", "INT", "PK", "รหัสอ้างอิงรายการย่อย (เพิ่มค่าอัตโนมัติ)"],
            ["order_id", "INT", "FK", "รหัสคำสั่งซื้อหลัก (เชื่อมโยงตาราง ORDERS)"],
            ["menu_id", "INT", "FK", "รหัสเมนูอาหาร (เชื่อมโยงตาราง MENUS)"],
            ["quantity", "INT", "-", "จำนวนรายการอาหารที่สั่ง"],
            ["unit_price", "DECIMAL(10,2)", "-", "ราคาต่อหน่วยของอาหาร ณ ขณะทำการสั่งซื้อ (บันทึกประวัติราคาคงที่)"],
            ["notes", "TEXT", "-", "หมายเหตุตัวเลือกพิเศษและความต้องการเพิ่มเติมจากลูกค้า (แสดงผลบนหน้าจอครัว KDS)"],
            ["created_at", "DATETIME", "-", "วันและเวลาที่บันทึกรายการอาหารย่อย (ค่าเริ่มต้นเวลาปัจจุบัน)"]
        ]
        while len(t_order_items.rows) > len(clean_order_items_data) + 1:
            t_order_items._tbl.remove(t_order_items.rows[-1]._tr)
        while len(t_order_items.rows) < len(clean_order_items_data) + 1:
            t_order_items.add_row()
        for i, row_data in enumerate(clean_order_items_data):
            set_table_row(t_order_items.rows[i + 1], row_data)

    # 4. Update INGREDIENTS table
    if t_ingredients:
        print("Updating INGREDIENTS table...")
        clean_ingredients_data = [
            ["ingredient_id", "INT", "PK", "รหัสอ้างอิงวัตถุดิบ (เพิ่มค่าอัตโนมัติ)"],
            ["name", "VARCHAR(100)", "-", "ชื่อวัตถุดิบ (เช่น มะละกอ, พริกสด, เนื้อไก่, หมูสด)"],
            ["quantity", "DECIMAL(10,2)", "-", "จำนวนวัตถุดิบที่เหลืออยู่ในสต็อกปัจจุบัน (รองรับการตัดสต็อกอัตโนมัติ)"],
            ["unit", "VARCHAR(50)", "-", "หน่วยนับของวัตถุดิบ (เช่น กิโลกรัม, กรัม, ลิตร)"],
            ["min_quantity", "DECIMAL(10,2)", "-", "จุดสั่งซื้อเพิ่ม / ปริมาณขั้นต่ำสำหรับแจ้งเตือนวัตถุดิบใกล้หมด"],
            ["updated_at", "DATETIME", "-", "วันและเวลาที่อัปเดตสต็อกล่าสุด"],
            ["created_at", "DATETIME", "-", "วันและเวลาที่สร้างรายการวัตถุดิบเข้าระบบ"]
        ]
        while len(t_ingredients.rows) > len(clean_ingredients_data) + 1:
            t_ingredients._tbl.remove(t_ingredients.rows[-1]._tr)
        while len(t_ingredients.rows) < len(clean_ingredients_data) + 1:
            t_ingredients.add_row()
        for i, row_data in enumerate(clean_ingredients_data):
            set_table_row(t_ingredients.rows[i + 1], row_data)

    # 5. Insert Database Views (Tables 13-17) and Database Indexes (Table 18) sequentially after Table 12
    if t_menu_ingredients:
        # Check if already added
        already_added = any("ORDER_SUMMARIES_VIEW" in p.text for p in doc.paragraphs)
        if not already_added:
            print("Inserting Database Views (3.4.4) and Indexes (3.4.5) sequentially after Table 12...")
            cur = t_menu_ingredients._element
            headers_4col = ["ชื่อฟิลด์ (Field)", "ชนิดข้อมูล (Data Type)", "คีย์ (Key)", "คำอธิบาย (Description)"]

            # Section 3.4.4 Heading
            p_v_head = make_heading_paragraph(doc, "3.4.4 การออกแบบมุมมองข้อมูล (Database Views)", bold=True, space_before=Pt(14), space_after=Pt(4))
            cur = insert_after(cur, p_v_head._element)

            p_v_desc = make_body_paragraph(doc, "เพื่อเพิ่มประสิทธิภาพในการสืบค้นข้อมูลและลดภาระการประมวลผลการคำนวณที่ซับซ้อน (Complex Query / Aggregations) ของระบบ คณะผู้จัดทำได้ออกแบบและสร้างมุมมองข้อมูล (Database Views) บนฐานข้อมูล MySQL จำนวน 5 รายการ เพื่อนำไปใช้งานร่วมกับแดชบอร์ดของผู้ดูแลระบบ รายงานสรุปยอดขาย การแจ้งเตือนสต็อกห้องครัว และการติดตามสถานะโต๊ะอาหารแบบเรียลไทม์ ดังนี้", space_after=Pt(6))
            cur = insert_after(cur, p_v_desc._element)

            # View 1: ORDER_SUMMARIES_VIEW
            p_t13_h = make_heading_paragraph(doc, "ตารางที่ 13: รายละเอียดมุมมองข้อมูลสรุปยอดขายรวม (ORDER_SUMMARIES_VIEW)", bold=True, space_before=Pt(6), space_after=Pt(4))
            cur = insert_after(cur, p_t13_h._element)
            p_t13_d = make_body_paragraph(doc, "ใช้สำหรับสรุปภาพรวมคำสั่งซื้อ ยอดเงิน และสถานะบิล สำหรับแสดงผลบนแดชบอร์ดผู้ดูแลระบบและรายงานสรุป", space_after=Pt(4))
            cur = insert_after(cur, p_t13_d._element)
            t13 = make_styled_table(doc, headers_4col, [
                ["order_id", "INT", "PK (จำลอง)", "รหัสอ้างอิงคำสั่งซื้อ"],
                ["table_number", "VARCHAR(10)", "-", "หมายเลขโต๊ะอาหารที่สั่ง (อนุญาตค่าว่างกรณีเดลิเวอรี)"],
                ["total_price", "DECIMAL(10,2)", "-", "ยอดเงินรวมสุทธิของคำสั่งซื้อ"],
                ["status", "VARCHAR(30)", "-", "สถานะปัจจุบันของคำสั่งซื้อ"],
                ["order_date", "DATETIME", "-", "วันและเวลาที่ทำการสั่งซื้ออาหาร"]
            ])
            cur = insert_after(cur, t13._element)
            p_space1 = make_body_paragraph(doc, "", space_after=Pt(6))
            cur = insert_after(cur, p_space1._element)

            # View 2: TRANSACTION_RECEIPTS_VIEW
            p_t14_h = make_heading_paragraph(doc, "ตารางที่ 14: รายละเอียดมุมมองข้อมูลใบเสร็จรับเงิน (TRANSACTION_RECEIPTS_VIEW)", bold=True, space_before=Pt(6), space_after=Pt(4))
            cur = insert_after(cur, p_t14_h._element)
            p_t14_d = make_body_paragraph(doc, "ใช้สำหรับตรวจสอบประวัติการทำธุรกรรมการเงินและออกใบเสร็จรับเงิน โดยรวมข้อมูลธุรกรรม คำสั่งซื้อ และชื่อลูกค้าเข้าด้วยกัน", space_after=Pt(4))
            cur = insert_after(cur, p_t14_d._element)
            t14 = make_styled_table(doc, headers_4col, [
                ["transaction_id", "INT", "PK (จำลอง)", "รหัสอ้างอิงธุรกรรมการชำระเงิน"],
                ["order_id", "INT", "FK", "รหัสอ้างอิงคำสั่งซื้อหลัก"],
                ["customer_name", "VARCHAR(100)", "-", "ชื่อลูกค้าหรือบัญชีผู้ใช้งานที่ทำรายการสั่งซื้อ"],
                ["total_amount", "DECIMAL(10,2)", "-", "จำนวนเงินรวมที่ชำระในธุรกรรมนี้"],
                ["payment_method", "VARCHAR(50)", "-", "ช่องทางหรือวิธีการชำระเงิน (เช่น PromptPay, Cash)"],
                ["payment_status", "VARCHAR(20)", "-", "สถานะการชำระเงิน (เช่น COMPLETED, PENDING)"],
                ["payment_date", "DATETIME", "-", "วันและเวลาที่บันทึกการทำธุรกรรมชำระเงิน"]
            ])
            cur = insert_after(cur, t14._element)
            p_space2 = make_body_paragraph(doc, "", space_after=Pt(6))
            cur = insert_after(cur, p_space2._element)

            # View 3: TOP_SELLING_MENUS_VIEW
            p_t15_h = make_heading_paragraph(doc, "ตารางที่ 15: รายละเอียดมุมมองข้อมูลเมนูอาหารยอดนิยม (TOP_SELLING_MENUS_VIEW)", bold=True, space_before=Pt(6), space_after=Pt(4))
            cur = insert_after(cur, p_t15_h._element)
            p_t15_d = make_body_paragraph(doc, "ใช้สำหรับวิเคราะห์และจัดอันดับเมนูขายดีประจำร้าน แสดงผลจำนวนจานที่ขายได้และรายได้รวมของแต่ละเมนู", space_after=Pt(4))
            cur = insert_after(cur, p_t15_d._element)
            t15 = make_styled_table(doc, headers_4col, [
                ["menu_id", "INT", "PK (จำลอง)", "รหัสอ้างอิงเมนูอาหาร"],
                ["menu_name", "VARCHAR(150)", "-", "ชื่อรายการอาหาร"],
                ["category_name", "VARCHAR(100)", "-", "ชื่อหมวดหมู่อาหารที่สังกัด"],
                ["total_sold", "INT", "-", "จำนวนจานที่ขายได้สะสมทั้งหมด"],
                ["total_revenue", "DECIMAL(10,2)", "-", "ยอดรายได้รวมสะสมที่ได้จากการขายเมนูนี้"]
            ])
            cur = insert_after(cur, t15._element)
            p_space3 = make_body_paragraph(doc, "", space_after=Pt(6))
            cur = insert_after(cur, p_space3._element)

            # View 4: LOW_STOCK_ALERTS_VIEW
            p_t16_h = make_heading_paragraph(doc, "ตารางที่ 16: รายละเอียดมุมมองข้อมูลแจ้งเตือนวัตถุดิบใกล้หมด (LOW_STOCK_ALERTS_VIEW)", bold=True, space_before=Pt(6), space_after=Pt(4))
            cur = insert_after(cur, p_t16_h._element)
            p_t16_d = make_body_paragraph(doc, "ใช้สำหรับแจ้งเตือนพนักงานและผู้จัดการแบบอัตโนมัติ เมื่อปริมาณวัตถุดิบคงเหลือต่ำกว่าหรือเท่ากับจุดสั่งซื้อเพิ่ม (Reorder Level)", space_after=Pt(4))
            cur = insert_after(cur, p_t16_d._element)
            t16 = make_styled_table(doc, headers_4col, [
                ["ingredient_id", "INT", "PK (จำลอง)", "รหัสอ้างอิงวัตถุดิบ"],
                ["ingredient_name", "VARCHAR(100)", "-", "ชื่อรายการวัตถุดิบ"],
                ["quantity_in_stock", "DECIMAL(10,2)", "-", "ปริมาณวัตถุดิบคงเหลือในคลังปัจจุบัน"],
                ["reorder_level", "DECIMAL(10,2)", "-", "จุดสั่งซื้อเพิ่ม (เกณฑ์ขั้นต่ำสำหรับแจ้งเตือนเตือนภัยสต็อก)"],
                ["unit", "VARCHAR(50)", "-", "หน่วยนับของวัตถุดิบ (เช่น กก., ลิตร)"]
            ])
            cur = insert_after(cur, t16._element)
            p_space4 = make_body_paragraph(doc, "", space_after=Pt(6))
            cur = insert_after(cur, p_space4._element)

            # View 5: LIVE_TABLE_STATUS_VIEW
            p_t17_h = make_heading_paragraph(doc, "ตารางที่ 17: รายละเอียดมุมมองข้อมูลสถานะโต๊ะอาหารแบบเรียลไทม์ (LIVE_TABLE_STATUS_VIEW)", bold=True, space_before=Pt(6), space_after=Pt(4))
            cur = insert_after(cur, p_t17_h._element)
            p_t17_d = make_body_paragraph(doc, "ใช้สำหรับหน้าจอติดตามสถานะโต๊ะอาหาร (Table Monitoring) แสดงสถานะโต๊ะ ออเดอร์ที่กำลังดำเนินงาน และยอดรวมบิลปัจจุบัน", space_after=Pt(4))
            cur = insert_after(cur, p_t17_d._element)
            t17 = make_styled_table(doc, headers_4col, [
                ["table_id", "INT", "PK (จำลอง)", "รหัสอ้างอิงโต๊ะอาหาร"],
                ["table_number", "VARCHAR(10)", "-", "หมายเลขโต๊ะอาหาร"],
                ["capacity", "INT", "-", "จำนวนที่นั่งรองรับสูงสุด"],
                ["current_status", "VARCHAR(20)", "-", "สถานะปัจจุบันของโต๊ะ (เช่น AVAILABLE, OCCUPIED)"],
                ["current_order_id", "INT", "-", "รหัสคำสั่งซื้อปัจจุบันที่กำลังรับประทานหรือรอชำระเงิน (ถ้ามี)"],
                ["current_total", "DECIMAL(10,2)", "-", "ยอดรวมค่าอาหารปัจจุบันของโต๊ะที่กำลังเปิดบิลอยู่"]
            ])
            cur = insert_after(cur, t17._element)
            p_space5 = make_body_paragraph(doc, "", space_after=Pt(8))
            cur = insert_after(cur, p_space5._element)

            # Section 3.4.5: Database Indexes
            p_idx_h = make_heading_paragraph(doc, "3.4.5 การออกแบบดัชนีข้อมูลเพื่อเพิ่มประสิทธิภาพการสืบค้น (Database Indexes)", bold=True, space_before=Pt(14), space_after=Pt(4))
            cur = insert_after(cur, p_idx_h._element)

            p_idx_d = make_body_paragraph(doc, "เพื่อเพิ่มประสิทธิภาพในการสืบค้นข้อมูล (Query Optimization) และรองรับการทำงานพร้อมกันของคำสั่งซื้อจำนวนมาก (High Concurrency) ระบบได้มีการกำหนดดัชนี (Indexes) บนคอลัมน์ที่มีการเรียกค้น กรองสถานะ และเชื่อมโยงความสัมพันธ์บ่อยครั้ง ดังแสดงในตารางที่ 18", space_after=Pt(6))
            cur = insert_after(cur, p_idx_d._element)

            p_t18_h = make_heading_paragraph(doc, "ตารางที่ 18: รายละเอียดการสร้างดัชนี (Indexes) เพื่อเพิ่มความเร็วของระบบฐานข้อมูล", bold=True, space_before=Pt(6), space_after=Pt(4))
            cur = insert_after(cur, p_t18_h._element)

            headers_index = ["ชื่อตาราง (Table)", "ชื่อดัชนี (Index Name)", "ฟิลด์ที่ทำดัชนี (Indexed Columns)", "ประเภท (Type)", "วัตถุประสงค์ (Purpose)"]
            index_rows = [
                ["USERS", "USERS_username_key", "username", "UNIQUE", "ป้องกันชื่อผู้ใช้งานซ้ำและเพิ่มความเร็วในการตรวจสอบ Login"],
                ["PROMOTIONS", "PROMOTIONS_code_key", "code", "UNIQUE", "ป้องกันรหัสคูปองซ้ำและค้นหารหัสส่วนลดได้อย่างรวดเร็ว"],
                ["MENUS", "MENUS_category_id_idx", "category_id", "INDEX", "เพิ่มความเร็วในการกรองเมนูตามหมวดหมู่บนหน้าจอสั่งอาหาร"],
                ["MENU_ALLERGENS", "MENU_ALLERGENS_menu_id_idx", "menu_id", "INDEX", "เพิ่มความเร็วในการเชื่อมโยงสารก่อภูมิแพ้ของแต่ละเมนู"],
                ["MENU_INGREDIENTS", "MENU_INGREDIENTS_menu_id_idx", "menu_id", "INDEX", "เพิ่มความเร็วในการคำนวณและตัดสต็อกวัตถุดิบตามสูตรอาหาร"],
                ["ORDERS", "ORDERS_user_id_idx", "user_id", "INDEX", "เพิ่มความเร็วในการดึงประวัติการสั่งซื้อของลูกค้าแต่ละราย"],
                ["ORDERS", "ORDERS_table_id_idx", "table_id", "INDEX", "เพิ่มความเร็วในการค้นหาคำสั่งซื้อประจำโต๊ะอาหาร"],
                ["ORDERS", "ORDERS_status_idx", "status", "INDEX", "เพิ่มความเร็วในการกรองสถานะออเดอร์ในหน้าจอครัว (KDS) และผู้ส่งอาหาร"],
                ["ORDERS", "ORDERS_created_at_idx", "created_at", "INDEX", "เพิ่มความเร็วในการจัดเรียงออเดอร์ล่าสุดและการออกรายงานยอดขายตามช่วงเวลา"],
                ["ORDER_ITEMS", "ORDER_ITEMS_order_id_idx", "order_id", "INDEX", "เพิ่มความเร็วในการดึงรายการอาหารย่อยภายใต้คำสั่งซื้อหลัก"],
                ["TRANSACTIONS", "TRANSACTIONS_order_id_idx", "order_id", "INDEX", "เพิ่มความเร็วในการค้นหาธุรกรรมตามรหัสคำสั่งซื้อ"],
                ["TRANSACTIONS", "TRANSACTIONS_payment_status_idx", "payment_status", "INDEX", "เพิ่มความเร็วในการตรวจสอบและกรองสถานะการชำระเงิน"]
            ]
            t18 = make_styled_table(doc, headers_index, index_rows)
            cur = insert_after(cur, t18._element)
            p_space6 = make_body_paragraph(doc, "", space_after=Pt(8))
            cur = insert_after(cur, p_space6._element)

    doc.save(doc_path)
    print("Document successfully updated and saved!")

if __name__ == "__main__":
    doc_path = "ระบบสั่งอาหารออนไลน์แบบครบวงจร-กรณีศึกษาร้านตำครกซิ่ง-ล่าสุด.docx"
    update_document(doc_path)
